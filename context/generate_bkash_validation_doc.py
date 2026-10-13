import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # Set margins to 1 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Style colors
    RED_COLOR = RGBColor(180, 0, 0)
    DARK_COLOR = RGBColor(30, 30, 30)

    def add_header(title, center=True, underline=True, size=16):
        p = doc.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(title)
        run.bold = True
        run.underline = underline
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.color.rgb = DARK_COLOR
        return p

    def add_api_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(14)
        run.font.color.rgb = DARK_COLOR
        return p

    def add_bullet_item(label, value=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.2)
        
        # Bullet dot
        r_bullet = p.add_run("• ")
        r_bullet.bold = True
        r_bullet.font.name = "Arial"
        r_bullet.font.size = Pt(10.5)
        r_bullet.font.color.rgb = RED_COLOR

        # Label
        r_label = p.add_run(f"{label} : ")
        r_label.bold = True
        r_label.font.name = "Arial"
        r_label.font.size = Pt(10.5)
        r_label.font.color.rgb = RED_COLOR

        # Value
        if value:
            r_val = p.add_run(value)
            r_val.font.name = "Arial"
            r_val.font.size = Pt(10)
            r_val.font.color.rgb = DARK_COLOR
        return p

    def add_json_block(json_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(json_text)
        run.font.name = "Courier New"
        run.font.size = Pt(9.5)
        run.font.color.rgb = DARK_COLOR
        return p

    # ================= PAGE 1 =================
    add_header("API Request/Response", center=True, underline=True, size=15)

    # 1. Grant Token
    add_api_section_header("Grant Token")
    add_bullet_item("API Title", "Grant Token")
    add_bullet_item("API URL", "https://tokenized.sandbox.bka.sh/v2/tokenized-checkout/auth/grant-token")
    add_bullet_item("Request Body")
    add_json_block('{\n  "app_key": "0vWQuCRGiUX7EPVjQDr0EUAYtc",\n  "app_secret": "jcUNPBgbcqEDedNKdvE4G1cAK7D3hCjmJccNPZZBq96QIxxwAMEx"\n}')
    add_bullet_item("API Response")
    add_json_block('{\n  "statusCode": "0000",\n  "statusMessage": "Successful",\n  "token_type": "Bearer",\n  "id_token": "eyJraWQiOiJvTVJzNU9ZY0wrUnRXQ2o3ZEJtdlc5VDBEcytrckw5M1NzY0VqUzlERXVzPSIsImFsZyI6IlJTMjU2In0.eyJzdWIiOiIzZmNmMTFhOC0yYTI0LTQ5YjAtYjZlNC1iYTZkOGFiOWNlNDgiLCJhdWQiOiI2cDdhcWVzZmljZTAxazltNWdxZTJhMGlhaCIsImV2ZW50X2lkIjoiZThlZGJmZmQtOGQ1OC00MTA4LWIxMDktY2MzYjZlNjYzYjNjIiwidG9rZW5fdXNlIjoiaWQiLCJhdXRoX3RpbWUiOjE3OTE4ODUwOTgsImlzcyI6Imh0dHBzOi8vY29nbml0by1pZHAuYXAtc291dGhlYXN0LTEuYW1hem9uYXdzLmNvbS9hcC1zb3V0aGVhc3QtMV9yYTNuUFkzSlMiLCJjb2duaXRvOnVzZXJuYW1lIjoiMDE3NzA2MTg1NjciLCJleHAiOjE3OTE4ODg2OTgsImlhdCI6MTc5MTg4NTA5OH0.Q321tkc8R8WPD6LisAA2oR7ieJZmZsze4Aq7wzL0bl37qVDFFgG0oxU9bnl1-DLO1cdvMgxv05ovJ-PCNopmc1dLlaWkB5kfnhmAnJPElX85EMgINh8agvGlGM0j3urCa54xMjZapjSpYg5BPswS9Wn5JUT45Q1_pq7uF02LI8V-kNYLZm4HmXccA725ZC8SJUsKM80t2xuJ7NI9wlUF2DsSZnQMUBvlxlLjsW_KkFS9oDNyj5CP7lambT_JYVW1oXiO31OUtL3cY7tzO8nRTqZmFKVcimKTViy7XayjTs2liBnMWOZgpN16RGNlofPA8N2aokxDQeM3fEARJO26-g",\n  "refresh_token": "eyJjdHkiOiJKV1QiLCJlbmMiOiJBMjU2R0NNIiwiYWxnIjoiUlNBLU9BRVAifQ.dv0_cWzfJHFpD_iDgk5K9PoI2LRxqQlG0smNVyz372ODRbP8knGU2jzCtfE3M5uy7ZSUHj-WETpID1Hfze0VYo8gloxyr3-BxwUJHCXZlm-veam62mquvvsY55s_2ciPNZUjoq2ZCHZDM91Aj-jI-H6XIaAr2m9BAnxLz6AxGq-fZvTK7ArdNcuHLEjGNdOwEZwZalHN-tDFzfvDeq3GIPCN-tVhfG3JV6Znj0n_SycEth3WmBEcpM8GruPeN3lxIT-U6DbNuRDri9lIi1bWZqOjJXrsWD1qx7fkFbNdc0pBVuKf5CsLGAFEOrgq3aZmhDSYgI3siEwP_jEFeznETg.-q31nd_NzU3MHcgx.RupgHt48C7WdHwKGkOnaiIWMaTmaPeHnYA7PvxSYK_BLTV0TAO5l9v4WPgihJTjEhsgY009EGzO-saID2R2UpXsp2Nq-mksC-fjzbgBozp3prhizm5LP5hCwXmtkczhlO-U3nMggmSjjfq5q0tNQ-bK_Mh_PVZ4qErEfzDk_YyP7nXvI49nyjvWyL0XpoMBlGYhX6MkF56V-aTq71p-Mi56MQbWojqZHM3ozxeyJ2mjhuI6xzGV0_IryJjYN5J1kgzwOYzw3rLhaLd4F-dY4d_XEXOvUVHequ0KB_LURmAHUGiLKU0gjmzl7z5PYZgj57L-0vlncvvPsKGym-U4wP0ktjSyv1KPhvbgtJAjtl25I0eL_qOtyDZTGUQ-Bn1UxFraWebgZzeoW4XHdmFAcze4bvbO5_cnvHpjYR7nQWYN55u9JuZecgYDVtThSXhgCT8Rguhs56Yzk5v1wlRBoZeEoJnycogLPSYZVz69KHjNysLtD7pfLmCsaZS2XqAlae0UtmYYw2N3wB63dmVzOyhE3IcRCE8_yyWccDikisYewv8AM0ZYxy1zv5cr80k4sazIEro5bBDR64yz-uE8iCFNQsi9Pg63B0jKOF6QWpuLOdINvp-32EZxseO9emwMfCSjDMhB0-nn07MBd096_jlgV_0hq4jYL6Au00C65X2WvnarCvEBNz-9QoyOKGek1JgprI4S-BrP0H7npSQN1a32HRg0Y4ojTcNhdz6kmd8KOnXQjq1WHGRiJ8q-5xJeRhc0fl9wAZn13m27PvfSBToe684CKdy36j6ShcwsqfXXa4GguBHPWHZpwyRZ3fFfIOhCrjjWo-SMsRtKvbnc52QuR_pUMSsMmPtuXK7CNHC1xUjnCs7uuq0_DkOU3MCLB2p0NW8mpAXpvNWlMf_FnGllAkFZFWn9IN0RWu0FBQUAJ2vW5_qzDd0ASyA632DQE1fjfWomGG3m8nGNvNw3RnTN2KGFLoq0-1GxDvw-az_sOebdeLQcFEh_MbxaSJLFqV8Zfs4MESuIiSsAY_zMwkrfEEiE5br9MIJnSUwzxTQiq0C2TYzm84VGan7Cx8TlpGMTxFg4JskSCZ7lgbYjbC-Lza8FPdVePmjPP_BFyuCpKcbAuTywJuz5-AifaHIIYPqCAypTGXJKqN5etJEKFM0qmUxwNyaIK17FbxgD6j9706cj6Df5PGE6kWhGkOpoCBE8PoWWp85-pzJ4ETAjMZA7SAqKnnctVVRU_S_tqioryCWzCt0Xe4FRIhsO5Zk4l0435tGVSgOASQTlTQOniku5WbA.kHMj09dC6416hKbroGYsMQ",\n  "expires_in": 3600\n}')

    # ================= PAGE 2 =================
    doc.add_page_break()

    # 2. Create API
    add_api_section_header("Create API")
    add_bullet_item("API Title", "Create Payment")
    add_bullet_item("API URL", "https://tokenized.sandbox.bka.sh/v2/tokenized-checkout/payment/create")
    add_bullet_item("Request Body")
    add_json_block('{\n  "mode": "0011",\n  "amount": "100",\n  "currency": "BDT",\n  "intent": "sale",\n  "merchantInvoiceNumber": "MINAN-1791885368271",\n  "callbackURL": "http://localhost:3001/api/bkash/callback",\n  "payerReference": "01770618575"\n}')
    add_bullet_item("API Response")
    add_json_block('{\n  "statusCode": "0000",\n  "statusMessage": "Successful",\n  "paymentID": "TR0011q1ohjxv1791885364885",\n  "bkashURL": "https://sandbox.payment.bkash.com/?paymentId=TR0011q1ohjxv1791885364885&hash=!.2((55mQiyK75d)!0(PJ71kK3oBJ1U2gegGhYd(OJZc09(1FjA)5ZXDYWr_f1I)DjVq(PjykSpvmn_.au.eCDa6oRv8W9qZxiOC1791885364885&mode=0011&apiVersion=v2/",\n  "callbackURL": "http://localhost:3001/api/bkash/callback",\n  "successCallbackURL": "http://localhost:3001/api/bkash/callback?paymentID=TR0011q1ohjxv1791885364885&status=success&signature=sGDVy310NN",\n  "failureCallbackURL": "http://localhost:3001/api/bkash/callback?paymentID=TR0011q1ohjxv1791885364885&status=failure&signature=sGDVy310NN",\n  "cancelledCallbackURL": "http://localhost:3001/api/bkash/callback?paymentID=TR0011q1ohjxv1791885364885&status=cancel&signature=sGDVy310NN",\n  "amount": "100",\n  "intent": "sale",\n  "currency": "BDT",\n  "paymentCreateTime": "2026-10-13T15:56:04:885 GMT+0600",\n  "transactionStatus": "Initiated",\n  "merchantInvoiceNumber": "MINAN-1791885368271",\n  "signature": "sGDVy310NN"\n}')

    # ================= PAGE 3 =================
    doc.add_page_break()

    # 3. Execute API
    add_api_section_header("Execute API")
    add_bullet_item("API Title", "Execute Payment")
    add_bullet_item("API URL", "https://tokenized.sandbox.bka.sh/v2/tokenized-checkout/payment/execute")
    add_bullet_item("Request Body")
    add_json_block('{\n  "paymentID": "TR0011q1ohjxv1791885364885"\n}')
    add_bullet_item("API Response")
    add_json_block('{\n  "statusCode": "0000",\n  "statusMessage": "Successful",\n  "paymentID": "TR0011q1ohjxv1791885364885",\n  "payerReference": "01770618575",\n  "customerMsisdn": "01770618575",\n  "trxID": "9I4001SA8D",\n  "amount": "100",\n  "transactionStatus": "Completed",\n  "paymentExecuteTime": "2026-10-13T15:58:05:983 GMT+0600",\n  "currency": "BDT",\n  "intent": "sale",\n  "merchantInvoiceNumber": "MINAN-1791885368271"\n}')

    # 4. Query Payment API
    add_api_section_header("Query Payment API")
    add_bullet_item("API Title", "Query Payment API")
    add_bullet_item("API URL", "https://tokenized.sandbox.bka.sh/v2/tokenized-checkout/query/payment")
    add_bullet_item("Request Body")
    add_json_block('{\n  "paymentID": "TR0011q1ohjxv1791885364885"\n}')
    add_bullet_item("API Response")
    add_json_block('{\n  "statusCode": "0000",\n  "statusMessage": "Successful",\n  "paymentID": "TR0011q1ohjxv1791885364885",\n  "mode": "0011",\n  "paymentCreateTime": "2026-10-13T15:56:04:885 GMT+0600",\n  "paymentExecuteTime": "2026-10-13T15:58:05:983 GMT+0600",\n  "amount": "100.00",\n  "currency": "BDT",\n  "intent": "sale",\n  "merchantInvoice": "MINAN-1791885368271",\n  "trxID": "9I4001SA8D",\n  "transactionStatus": "Completed",\n  "verificationStatus": "Complete",\n  "payerReference": "01770618575"\n}')

    # ================= PAGE 4 =================
    doc.add_page_break()

    # 5. Search Transaction API
    add_api_section_header("Search Transaction API")
    add_bullet_item("API Title", "Search Payment API")
    add_bullet_item("API URL", "https://tokenized.sandbox.bka.sh/v2/tokenized-checkout/general/search-transaction")
    add_bullet_item("Request Body")
    add_json_block('{\n  "trxID": "9I4001SA8D"\n}')
    add_bullet_item("API Response")
    add_json_block('{\n  "statusCode": "2003",\n  "statusMessage": "Process Failed"\n}')

    # ================= PAGE 5 =================
    doc.add_page_break()

    add_header("Error Message Implementation", center=True, underline=True, size=15)

    def add_case_meta(case_title, inv, pid, time_str):
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(10)
        p_title.paragraph_format.space_after = Pt(4)
        r = p_title.add_run(case_title)
        r.bold = True
        r.underline = True
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.color.rgb = DARK_COLOR

        p_inv = doc.add_paragraph()
        p_inv.paragraph_format.space_before = Pt(2)
        p_inv.paragraph_format.space_after = Pt(2)
        r1 = p_inv.add_run("Invoice number: ")
        r1.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(10.5)
        p_inv.add_run(inv).font.size = Pt(10.5)

        p_pid = doc.add_paragraph()
        p_pid.paragraph_format.space_before = Pt(2)
        p_pid.paragraph_format.space_after = Pt(2)
        r_pid = p_pid.add_run("Payment ID: ")
        r_pid.bold = True
        r_pid.font.name = "Arial"
        r_pid.font.size = Pt(10.5)
        p_pid.add_run(pid).font.size = Pt(10.5)

        p_time = doc.add_paragraph()
        p_time.paragraph_format.space_before = Pt(2)
        p_time.paragraph_format.space_after = Pt(2)
        r2 = p_time.add_run("Time of Transaction: ")
        r2.bold = True
        r2.font.name = "Arial"
        r2.font.size = Pt(10.5)
        p_time.add_run(time_str).font.size = Pt(10.5)

        p_sc = doc.add_paragraph()
        p_sc.paragraph_format.space_before = Pt(2)
        p_sc.paragraph_format.space_after = Pt(6)
        r3 = p_sc.add_run("Screenshot:")
        r3.bold = True
        r3.font.name = "Arial"
        r3.font.size = Pt(10.5)

    def add_image_bordered(img_path):
        if not os.path.exists(img_path):
            print(f"Image not found: {img_path}")
            return
        
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(5.8)

        # Set thin gray border around cell
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
            f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
            f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(4)
        run = cp.add_run()
        run.add_picture(img_path, width=Inches(5.5))
        
        # Spacing after table
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after = Pt(8)

    # Case A
    add_case_meta(
        "Case#A (Duplicate Transaction)",
        "MN-20261013-0005-01",
        "TR00116Zr5hJZ1791890986014",
        "05:30pm Oct 13 2026"
    )
    add_image_bordered(r"c:\new projects\minan\context\screenshots\case_a_duplicate_transaction.png")

    # Case B
    add_case_meta(
        "Case#B (Cancel Payment)",
        "MN-20261013-0003-01",
        "TR0011ur46kzL1791886294584",
        "04:12pm Oct 13 2026"
    )
    add_image_bordered(r"c:\new projects\minan\context\screenshots\case_b_cancel_payment.png")

    # ================= PAGE 6 =================
    doc.add_page_break()

    # Case C
    add_case_meta(
        "Case#C (Wrong OTP Exceed)",
        "MN-20261013-0004-01",
        "TR0011RFrPPVu1791888643661",
        "04:51pm Oct 13 2026"
    )
    add_image_bordered(r"c:\new projects\minan\context\screenshots\case_c_wrong_otp_exceed.png")

    out_path = r"c:\new projects\minan\context\bKash_Sandbox_Validation_Response.docx"
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == "__main__":
    create_document()
